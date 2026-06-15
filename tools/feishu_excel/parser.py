"""
Excel 文件解析器。
从 feishu-bot/main.py 提取的核心解析逻辑。
"""
import logging
import re
import urllib.parse
from io import BytesIO
from pathlib import Path

import openpyxl
import requests

log = logging.getLogger(__name__)

DATA_SHEET = "综拓开场数据基础模板2"
FALLBACK_SHEET_KEYWORDS = ["综拓", "开场数据", "基础模板"]

REQUIRED_COLUMNS = [
    "场次数",
    "劣势影城数",
    "排片占比",
]

OPTIONAL_COLUMNS = [
    "影片距离满足红包场次数差1场影城数",
    "影片距离满足红包场次数差2场影城数",
]

ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


# ---- 文件名解析 ----

def parse_filename(filename: str) -> dict | None:
    """从文件名解析影片名、监控日期、提取时间。"""
    name = Path(filename).stem
    t = r"(\d{2})[:,_](\d{2})(?:[:,_](\d{2}))?"
    pattern = (
        r"^(.+?)"
        r"\((\d{4}-\d{2}-\d{2})[+ ]" + t
        + r"-(\d{4}-\d{2}-\d{2})[+ ]" + t + r"\)"
        r"(\d{4}-\d{2}-\d{2})[+ ](\d{2})[:,_](\d{2})"
        r"$"
    )
    m = re.match(pattern, name)
    if not m:
        return None

    mon_start_str = f"{m.group(2)} {m.group(3)}:{m.group(4)}:{m.group(5) or '00'}"
    mon_end_str = f"{m.group(6)} {m.group(7)}:{m.group(8)}:{m.group(9) or '00'}"

    return {
        "movie": m.group(1),
        "mon_start": mon_start_str,
        "mon_end": mon_end_str,
        "extract_date": m.group(10),
        "extract_hour": int(m.group(11)),
    }


# ---- Excel 数据提取 ----

def find_data_sheet(wb) -> str | None:
    """找到包含目标列头的数据 Sheet。"""
    if DATA_SHEET in wb.sheetnames:
        return DATA_SHEET

    for name in wb.sheetnames:
        for kw in FALLBACK_SHEET_KEYWORDS:
            if kw in name:
                return name

    for name in wb.sheetnames:
        ws = wb[name]
        for row_idx in range(1, min(ws.max_row + 1, 10)):
            for c in range(1, min(ws.max_column + 1, 30)):
                val = str(ws.cell(row=row_idx, column=c).value or "")
                if "场次数" in val:
                    return name

    return None


def find_header_row(ws) -> int | None:
    for row_idx in range(1, min(ws.max_row + 1, 15)):
        for c in range(1, ws.max_column + 1):
            if "场次数" in str(ws.cell(row=row_idx, column=c).value or ""):
                return row_idx
    return None


def find_heji_row(ws, start_row: int) -> int | None:
    substring_match = None
    for row_idx in range(start_row + 1, ws.max_row + 1):
        val = str(ws.cell(row=row_idx, column=1).value or "").strip()
        if val == "合计":
            return row_idx
        if substring_match is None and "合计" in val:
            substring_match = row_idx
    return substring_match


def extract_all_data(ws, header_row: int, main_movie: str = "") -> dict | None:
    """提取主电影和对比电影（如有）的数据。"""
    heji_row = find_heji_row(ws, header_row)
    if heji_row is None:
        log.warning("未找到合计行")
        return None

    col_occurrences: dict[str, list[int]] = {}
    for col_idx in range(1, ws.max_column + 1):
        val = str(ws.cell(row=header_row, column=col_idx).value or "").strip()
        if val:
            if val not in col_occurrences:
                col_occurrences[val] = []
            col_occurrences[val].append(col_idx)

    is_cmp = "场次数" in col_occurrences and len(col_occurrences["场次数"]) >= 2

    main_col_map = {}
    for col_name in ALL_COLUMNS:
        if col_name in col_occurrences:
            main_col_map[col_name] = col_occurrences[col_name][0]

    missing = [c for c in REQUIRED_COLUMNS if c not in main_col_map]
    if missing:
        log.warning(f"缺少必要列: {missing}")
        return None

    main_data = {name: ws.cell(row=heji_row, column=col).value for name, col in main_col_map.items()}

    result = {"data": main_data, "cmp_movie": None, "cmp_data": None}

    if not is_cmp:
        return result

    first_cmp_col = col_occurrences["场次数"][1]
    cmp_movie = None
    for r in range(header_row - 1, max(0, header_row - 4), -1):
        val = str(ws.cell(row=r, column=first_cmp_col).value or "").strip()
        if val and val not in ("", "None") and val != main_movie and "场次" not in val:
            cmp_movie = val
            break

    if not cmp_movie:
        for r in range(header_row - 1, max(0, header_row - 4), -1):
            for c in range(first_cmp_col, ws.max_column + 1):
                val = str(ws.cell(row=r, column=c).value or "").strip()
                if val and len(val) > 1 and val != main_movie and "场次" not in val and "占比" not in val and "差值" not in val:
                    cmp_movie = val
                    break
            if cmp_movie:
                break

    if not cmp_movie:
        cmp_movie = "对比影片"

    log.info(f"对比模式: {main_movie} vs {cmp_movie}")

    cmp_col_map = {}
    cmp_shared = ["场次数", "排片占比", "劣势影城数"]
    for name in cmp_shared:
        if name in col_occurrences and len(col_occurrences[name]) >= 2:
            cmp_col_map[name] = col_occurrences[name][1]

    cmp_specific = [
        "场次数新增",
        "主影片与当前影片场次差值",
        "主影片与当前影片排片占比差值",
        "拍片占比",
    ]
    for name in cmp_specific:
        if name in col_occurrences:
            cmp_col_map[name] = col_occurrences[name][0]

    cmp_data = {name: ws.cell(row=heji_row, column=col).value for name, col in cmp_col_map.items()}

    if "拍片占比" in cmp_data and "排片占比" not in cmp_data:
        cmp_data["排片占比"] = cmp_data.pop("拍片占比")

    result["cmp_movie"] = cmp_movie
    result["cmp_data"] = cmp_data
    return result


# ---- 单文件解析入口 ----

def parse_single_file(fname: str, content: bytes) -> dict | None:
    """解析单个 Excel 文件，返回 entry dict 或 None。"""
    safe_name = re.sub(r'[<>:"/\\|?*]', '_', fname)
    log.info(f"文件: {safe_name}")

    info = parse_filename(safe_name)
    if info is None:
        log.warning(f"文件名格式不匹配: {safe_name}")
        return None

    log.info(f"影片: {info['movie']}  监控: {info['mon_start']} ~ {info['mon_end']}")

    try:
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    except Exception as e:
        log.warning(f"Excel 文件打开失败: {e}")
        return None

    sheet_name = find_data_sheet(wb)
    if sheet_name is None:
        log.warning(f"未找到数据 Sheet，可用: {wb.sheetnames}")
        wb.close()
        return None

    ws = wb[sheet_name]
    log.info(f"Sheet: {sheet_name}  (max_row={ws.max_row}, max_col={ws.max_column})")
    header_row = find_header_row(ws)
    if header_row is None:
        log.warning("未找到表头")
        wb.close()
        return None

    log.info(f"表头行: {header_row}")

    all_data = extract_all_data(ws, header_row, info["movie"])
    wb.close()

    if all_data is None:
        log.warning("数据提取失败")
        return None

    info["data"] = all_data["data"]
    if all_data["cmp_movie"]:
        info["cmp_movie"] = all_data["cmp_movie"]
        info["cmp_data"] = all_data["cmp_data"]

    log.info(f"数据: {info['data']}")
    if info.get("cmp_movie"):
        log.info(f"对比: {info['cmp_movie']} -> {info['cmp_data']}")
    return info


# ---- 文件下载 ----

def download_file(url: str) -> tuple[str | None, bytes | None]:
    """下载 Excel，返回 (文件名, 内容)。"""
    try:
        resp = requests.get(url, timeout=60, headers={"Accept-Encoding": "identity"})
        resp.raise_for_status()

        fname = None
        cd = resp.headers.get("Content-Disposition", "")
        if "filename*=" in cd:
            parts = cd.split("filename*=")
            if len(parts) > 1:
                encoded = parts[1].split(";")[0].strip()
                if "''" in encoded:
                    _, fname_encoded = encoded.split("''", 1)
                    fname = urllib.parse.unquote(fname_encoded)
        if not fname and "filename=" in cd:
            raw = cd.split("filename=")[1].split(";")[0].strip().strip('"')
            try:
                fname = raw.encode("latin-1").decode("utf-8")
            except (UnicodeDecodeError, UnicodeEncodeError):
                fname = raw
        if not fname:
            fname = url.split("/")[-1].split("?")[0]

        return fname, resp.content
    except Exception as e:
        log.error(f"下载失败: {e}")
        return None, None
