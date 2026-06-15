"""
预测 Excel 生成器。
原位于 Prediction-Bot/excel/generator.py，已内嵌到 Agent 项目中。
"""
import io
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter


HEADERS = [
    "日期",           # A
    "影片名称",       # B
    "累计新增占比",   # C
    "大盘场次",       # D
    "目前大盘场次",   # E
    "剩余场次",       # F  =D-E
    "剩余可开场次",   # G  =F*C
    "目前场次",       # H  (排片场次)
    "影片总场次",     # I  =G+H
    "落位占比",       # J  =I/D
]

COL_WIDTHS = [8, 22, 14, 14, 14, 14, 14, 14, 14, 14]

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BODY_FONT = Font(size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def generate_excel(
    date_str: str,
    movies: list[dict],
    dapan_total: int,
    total_show_count: int,
) -> io.BytesIO:
    """
    生成预测 Excel 文件。

    参数:
      date_str:         预测日期，如 "6.15"
      movies:           影片列表，每个 dict 包含:
                          - name: 影片名称
                          - cumulative_share: 累计新增占比 (C列, 小数如 0.176)
                          - show_count: 目前场次/排片场次 (H列)
      dapan_total:      大盘场次 (D列, 用户提供)
      total_show_count: 目前大盘场次 (E列, API 获取)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "影片落位预测"

    _write_headers(ws)
    _write_data(ws, date_str, movies, dapan_total, total_show_count)
    _format_sheet(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _write_headers(ws):
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


def _write_data(ws, date_str, movies, dapan_total, total_show_count):
    pct_fmt = '0.0%'
    num_fmt = '#,##0'

    for i, movie in enumerate(movies):
        row = i + 2
        name = movie["name"]
        cum_share = movie["cumulative_share"]
        show_count = movie.get("show_count", 0)

        ws.cell(row=row, column=1, value=date_str).border = THIN_BORDER                       # A: 日期

        cell_b = ws.cell(row=row, column=2, value=name)                                        # B: 影片名称
        cell_b.border = THIN_BORDER

        cell_c = ws.cell(row=row, column=3, value=cum_share)                                   # C: 累计新增占比
        cell_c.number_format = pct_fmt
        cell_c.border = THIN_BORDER

        cell_d = ws.cell(row=row, column=4, value=dapan_total)                                 # D: 大盘场次
        cell_d.number_format = num_fmt
        cell_d.border = THIN_BORDER

        cell_e = ws.cell(row=row, column=5, value=total_show_count)                            # E: 目前大盘场次
        cell_e.number_format = num_fmt
        cell_e.border = THIN_BORDER

        cell_f = ws.cell(row=row, column=6, value=f"=D{row}-E{row}")                           # F: 剩余场次
        cell_f.number_format = num_fmt
        cell_f.border = THIN_BORDER

        cell_g = ws.cell(row=row, column=7, value=f"=F{row}*C{row}")                           # G: 剩余可开场次
        cell_g.number_format = num_fmt
        cell_g.border = THIN_BORDER

        cell_h = ws.cell(row=row, column=8, value=show_count)                                  # H: 目前场次
        cell_h.number_format = num_fmt
        cell_h.border = THIN_BORDER

        cell_i = ws.cell(row=row, column=9, value=f"=G{row}+H{row}")                           # I: 影片总场次
        cell_i.number_format = num_fmt
        cell_i.border = THIN_BORDER

        cell_j = ws.cell(row=row, column=10, value=f"=ROUNDDOWN(I{row}/D{row},3)")                          # J: 落位占比
        cell_j.number_format = pct_fmt
        cell_j.border = THIN_BORDER


def _format_sheet(ws):
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.sheet_properties.tabColor = "4472C4"
    ws.freeze_panes = "A2"
